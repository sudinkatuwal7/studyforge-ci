from io import BytesIO

from django.test import SimpleTestCase
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from processing_service.chunking import chunk_pages
from processing_service.document_parser import DOCX_MIME, PPTX_MIME, extract_document
from processing_service.errors import PermanentProcessingError


class OfficeExtractionTests(SimpleTestCase):
    def test_docx_preserves_heading_paragraph_and_table_order_without_pages(self):
        document = Document()
        document.add_heading("Course outline", level=2)
        document.add_paragraph("First paragraph")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Name"
        table.cell(0, 1).text = "Meaning"
        document.add_paragraph("After the table")
        stream = BytesIO()
        document.save(stream)

        result = extract_document(stream.getvalue(), DOCX_MIME)
        self.assertIsNone(result.page_count)
        self.assertEqual(len(result.pages), 1)
        self.assertEqual(result.pages[0].page_number, 1)
        self.assertEqual(result.pages[0].text.splitlines(), [
            "[Heading 2] Course outline", "First paragraph",
            "[Table row] Name | Meaning", "After the table",
        ])
        self.assertEqual(result.pages[0].context_prefix, "[Heading 2] Course outline")
        self.assertEqual(chunk_pages(result.pages)[0].content, result.pages[0].text)
        self.assertEqual(chunk_pages(result.pages)[0].page_number, 1)

    def test_long_docx_sections_repeat_only_their_current_heading(self):
        document = Document()
        document.add_heading("First topic", level=1)
        document.add_paragraph("alpha " * 350)
        document.add_heading("Second topic", level=2)
        document.add_paragraph("beta " * 350)
        stream = BytesIO()
        document.save(stream)

        units = extract_document(stream.getvalue(), DOCX_MIME).pages
        self.assertEqual([unit.context_prefix for unit in units], [
            "[Heading 1] First topic", "[Heading 2] Second topic",
        ])
        chunks = chunk_pages(units)
        first = [chunk for chunk in chunks if "alpha" in chunk.content]
        second = [chunk for chunk in chunks if "beta" in chunk.content]
        self.assertGreater(len(first), 1)
        self.assertGreater(len(second), 1)
        self.assertTrue(all(chunk.content.startswith("[Heading 1] First topic\n") for chunk in first))
        self.assertTrue(all(chunk.content.startswith("[Heading 2] Second topic\n") for chunk in second))
        self.assertTrue(all(len(chunk.content) <= 1000 for chunk in chunks))

    def test_pptx_preserves_slide_titles_text_and_location(self):
        presentation = Presentation()
        for title, body in (("Overview", "First slide"), ("Details", "Second slide")):
            slide = presentation.slides.add_slide(presentation.slide_layouts[5])
            slide.shapes.title.text = title
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = body
            slide.notes_slide.notes_text_frame.text = "Private speaker note"
        stream = BytesIO()
        presentation.save(stream)

        result = extract_document(stream.getvalue(), PPTX_MIME)
        self.assertEqual(result.page_count, 2)
        self.assertEqual([unit.page_number for unit in result.pages], [1, 2])
        self.assertIn("First slide", result.pages[0].text)
        self.assertIn("[Slide 1: Overview]", result.pages[0].text)
        self.assertEqual(result.pages[0].context_prefix, "[Slide 1: Overview]")
        self.assertEqual(chunk_pages(result.pages)[0].content, result.pages[0].text)
        self.assertIn("Second slide", result.pages[1].text)
        self.assertNotIn("Private speaker note", "\n".join(unit.text for unit in result.pages))
        self.assertEqual([chunk.page_number for chunk in chunk_pages(result.pages)], [1, 2])

    def test_long_pptx_slide_repeats_its_title_on_every_chunk(self):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "Overview"
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "slide content " * 250
        stream = BytesIO()
        presentation.save(stream)

        units = extract_document(stream.getvalue(), PPTX_MIME).pages
        chunks = chunk_pages(units)
        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(chunk.content.startswith("[Slide 1: Overview]\n") for chunk in chunks))
        self.assertTrue(all(chunk.page_number == 1 and len(chunk.content) <= 1000 for chunk in chunks))

    def test_blank_documents_have_no_usable_chunks(self):
        document = Document()
        stream = BytesIO()
        document.save(stream)
        self.assertEqual(chunk_pages(extract_document(stream.getvalue(), DOCX_MIME).pages), [])
        presentation = Presentation()
        presentation.slides.add_slide(presentation.slide_layouts[6])
        stream = BytesIO()
        presentation.save(stream)
        self.assertEqual(chunk_pages(extract_document(stream.getvalue(), PPTX_MIME).pages), [])

    def test_parser_failure_is_permanent(self):
        for mime in (DOCX_MIME, PPTX_MIME):
            with self.subTest(mime=mime), self.assertRaises(PermanentProcessingError):
                extract_document(b"broken package", mime)
