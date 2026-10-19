"""Normalize PDF, DOCX, and PPTX text into ordered location-bearing units."""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO

from .errors import PermanentProcessingError
from .pdf_parser import extract_pdf_text

PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


@dataclass(frozen=True)
class TextUnit:
    # page_number is the persisted legacy location field: a PDF page, PPTX slide,
    # or DOCX's required internal compatibility value (never a page citation).
    page_number: int
    text: str


@dataclass(frozen=True)
class DocumentExtraction:
    pages: list[TextUnit]
    page_count: int | None
    likely_scanned: bool = False


def extract_document(content: bytes, mime_type: str) -> DocumentExtraction:
    """Dispatch extraction once, then return the common chunking input."""
    if mime_type == PDF_MIME:
        result = extract_pdf_text(BytesIO(content))
        return DocumentExtraction(
            [TextUnit(page.page_number, page.text) for page in result.pages],
            result.page_count,
            result.likely_scanned,
        )
    try:
        if mime_type == DOCX_MIME:
            return _extract_docx(content)
        if mime_type == PPTX_MIME:
            return _extract_pptx(content)
    except Exception as exc:
        raise PermanentProcessingError("Could not extract text from the Office document.") from exc
    raise PermanentProcessingError("Unsupported source MIME type.")


def _extract_docx(content: bytes) -> DocumentExtraction:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(BytesIO(content))
    lines = []
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            value = block.text.strip()
            if value:
                style_name = block.style.name or ""
                if style_name.lower().startswith("heading"):
                    level = style_name[len("Heading"):].strip()
                    value = f"[Heading {level or '1'}] {value}"
                lines.append(value)
        elif isinstance(block, Table):
            for row in block.rows:
                cells = [" ".join(p.text.strip() for p in cell.paragraphs if p.text.strip()) for cell in row.cells]
                if any(cells):
                    lines.append("[Table row] " + " | ".join(cells))
    # DOCX is unpaginated. The positive chunk location is schema compatibility only.
    return DocumentExtraction([TextUnit(1, "\n".join(lines))], None)


def _extract_pptx(content: bytes) -> DocumentExtraction:
    from pptx import Presentation

    presentation = Presentation(BytesIO(content))
    units = []
    for number, slide in enumerate(presentation.slides, 1):
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape and title_shape.has_text_frame else ""
        lines = [f"[Slide {number}{': ' + title if title else ''}]"]

        def add_shapes(shapes):
            for shape in shapes:
                if shape == title_shape:
                    continue
                if shape.has_text_frame:
                    value = shape.text.strip()
                    if value:
                        lines.append(value)
                if shape.shape_type == 6:  # grouped shapes
                    add_shapes(shape.shapes)

        add_shapes(slide.shapes)
        # A title alone is usable text; the marker alone is not.
        units.append(TextUnit(number, "\n".join(lines) if title or len(lines) > 1 else ""))
    return DocumentExtraction(units, len(units))
