"""Upload-time source validation; text extraction belongs to the worker."""

import hashlib
import re
import unicodedata
from pathlib import PurePosixPath
from zipfile import ZipFile

from pypdf import PdfReader

from apps.documents.errors import DocumentAPIError


MAX_PDF_BYTES = 20 * 1024 * 1024
MAX_DOCUMENT_BYTES = MAX_PDF_BYTES
MAX_ARCHIVE_ENTRIES = 2000
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_MEMBER_BYTES = 30 * 1024 * 1024
FORMATS = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}


def validate_document(upload):
    if not upload.size:
        raise DocumentAPIError("EMPTY_FILE")
    if upload.size > MAX_DOCUMENT_BYTES:
        raise DocumentAPIError("FILE_TOO_LARGE", 413)
    name = (upload.name or "").replace("\\", "/").rsplit("/", 1)[-1]
    name = "".join(char for char in name if not unicodedata.category(char).startswith("C"))
    suffix = PurePosixPath(name.lower()).suffix
    if suffix not in FORMATS or not name[:-len(suffix)].strip():
        raise DocumentAPIError("INVALID_FORMAT")
    name = name[:255 - len(suffix)] + name[-len(suffix):] if len(name) > 255 else name
    if suffix == ".pdf":
        return (*validate_pdf(upload), suffix)

    digest = hashlib.sha256()
    size = 0
    for chunk in upload.chunks():
        size += len(chunk)
        if size > MAX_DOCUMENT_BYTES:
            raise DocumentAPIError("FILE_TOO_LARGE", 413)
        digest.update(chunk)
    try:
        upload.seek(0)
        with ZipFile(upload) as archive:
            entries = archive.infolist()
            if not entries or len(entries) > MAX_ARCHIVE_ENTRIES:
                raise ValueError("Archive entry count is invalid")
            total = 0
            for entry in entries:
                total += entry.file_size
                if (
                    entry.flag_bits & 1
                    or entry.file_size > MAX_MEMBER_BYTES
                    or total > MAX_UNCOMPRESSED_BYTES
                    or (entry.compress_size == 0 and entry.file_size)
                    or (entry.compress_size and entry.file_size > entry.compress_size * 200)
                    or entry.filename.startswith("/")
                    or ".." in PurePosixPath(entry.filename).parts
                ):
                    raise ValueError("Unsafe Office package")
            names = set(archive.namelist())
            if len(names) != len(entries):
                raise ValueError("Duplicate Office package entries")
            main_part = "word/document.xml" if suffix == ".docx" else "ppt/presentation.xml"
            if not {"[Content_Types].xml", "_rels/.rels", main_part} <= names:
                raise ValueError("Wrong Office package type")
            if archive.testzip() is not None:
                raise ValueError("Corrupt Office package")
        upload.seek(0)
        if suffix == ".docx":
            from docx import Document
            document = Document(upload)
            # Force lazy XML/table access while validating, without extracting text.
            list(document.element.body.iter())
        else:
            from pptx import Presentation
            presentation = Presentation(upload)
            for slide in presentation.slides:
                list(slide.shapes)
    except Exception as exc:
        raise DocumentAPIError("INVALID_OFFICE") from exc
    finally:
        upload.seek(0)
    return name, digest.hexdigest(), suffix


def validate_pdf(upload):
    if not upload.size:
        raise DocumentAPIError("EMPTY_FILE")
    if upload.size > MAX_PDF_BYTES:
        raise DocumentAPIError("FILE_TOO_LARGE", 413)

    name = (upload.name or "").replace("\\", "/").rsplit("/", 1)[-1]
    name = "".join(char for char in name if not unicodedata.category(char).startswith("C"))
    if not name.lower().endswith(".pdf") or not name[:-4].strip():
        raise DocumentAPIError("INVALID_PDF")
    name = name[:251] + name[-4:] if len(name) > 255 else name

    digest = hashlib.sha256()
    size = 0
    for chunk in upload.chunks():
        size += len(chunk)
        if size > MAX_PDF_BYTES:
            raise DocumentAPIError("FILE_TOO_LARGE", 413)
        digest.update(chunk)
    if not size:
        raise DocumentAPIError("EMPTY_FILE")

    try:
        upload.seek(0)
        if not re.match(rb"%PDF-(?:1\.[0-7]|2\.0)(?:\r|\n)", upload.read(16)):
            raise DocumentAPIError("INVALID_PDF")
        upload.seek(0)
        reader = PdfReader(upload, strict=True)
        if reader.is_encrypted:
            raise DocumentAPIError("UNSUPPORTED_PDF")
        if not len(reader.pages):
            raise DocumentAPIError("INVALID_PDF")
        # Force lazy page references and required geometry to be read. Image-only
        # and blank pages remain valid; text/content extraction belongs to the worker.
        for page in reader.pages:
            if page.get("/Type") != "/Page":
                raise DocumentAPIError("INVALID_PDF")
            box = page.mediabox
            if len(box) != 4:
                raise DocumentAPIError("INVALID_PDF")
            for coordinate in box:
                float(coordinate)
    except DocumentAPIError:
        raise
    except Exception as exc:
        # Parser failures must not leak document content or parser internals.
        raise DocumentAPIError("INVALID_PDF") from exc
    finally:
        upload.seek(0)

    return name, digest.hexdigest()
